from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output' / 'documents'
OUT.mkdir(parents=True, exist_ok=True)
source = Path(r'C:/Users/user/Downloads/Log Book DIPLOMA BARUUU (1).docx')
with ZipFile(source) as z:
    src = etree.fromstring(z.read('word/document.xml'))
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
body = src.find('w:body', ns)
def original(i):
    return ''.join(body[i].xpath('.//w:t/text()', namespaces=ns)).strip()

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.75)
section.left_margin = section.right_margin = Inches(.85)
section.footer_distance = Inches(.3)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1
if 'Placeholder Text' not in doc.styles:
    placeholder_style = doc.styles.add_style('Placeholder Text', WD_STYLE_TYPE.CHARACTER)
else:
    placeholder_style = doc.styles['Placeholder Text']
placeholder_style.font.color.rgb = RGBColor.from_string('777777')
for border in doc.styles['Title']._element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
for name, size in [('Title', 24), ('Heading 1', 14), ('Heading 2', 11)]:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(10)
    style.paragraph_format.space_after = Pt(6)

field_id = 100
def field(p, label='Klik di sini untuk isi penerangan.'):
    global field_id
    field_id += 1
    sdt = OxmlElement('w:sdt')
    pr = OxmlElement('w:sdtPr')
    for tag, val in [('alias', label), ('tag', f'ruang_{field_id}'), ('id', str(field_id))]:
        e = OxmlElement('w:' + tag); e.set(qn('w:val'), val); pr.append(e)
    pr.append(OxmlElement('w:showingPlcHdr'))
    sdt.append(pr)
    content = OxmlElement('w:sdtContent')
    r = OxmlElement('w:r'); rp = OxmlElement('w:rPr')
    charstyle = OxmlElement('w:rStyle'); charstyle.set(qn('w:val'), placeholder_style.style_id); rp.append(charstyle)
    r.append(rp); t = OxmlElement('w:t'); t.text = label; r.append(t)
    content.append(r); sdt.append(content); p._p.append(sdt)

def table(rows, cols, widths):
    t = doc.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        edge = OxmlElement('w:' + name)
        for attr, val in [('val', 'single'), ('sz', '4'), ('color', 'B7B7B7')]: edge.set(qn('w:' + attr), val)
        borders.append(edge)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for name in ['top', 'left', 'bottom', 'right']:
        e = OxmlElement('w:' + name); e.set(qn('w:w'), '110'); e.set(qn('w:type'), 'dxa'); margins.append(e)
    props.append(margins)
    for row in t.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    return t

def newpage(title):
    p = doc.add_paragraph(title, 'Heading 1')
    p.paragraph_format.page_break_before = True
    return p

p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(65)
p.add_run().add_picture(str(Path(__file__).parent / 'image1.png'), width=Inches(3.6))
for text, style in [('BUKU LOG', 'Title'), ('LATIHAN INDUSTRI', 'Title'), ('ITD 30110', 'Heading 1'), ('SEM I SESI 2026/2027', 'Heading 1'), ('FAKULTI INFORMATIK DAN KOMPUTERAN', 'Heading 1')]:
    p = doc.add_paragraph(text, style); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)

newpage('MAKLUMAT PELAJAR')
details = [
 ('Nama', 'MUHAMMAD ALIF BIN YAHYA'), ('No. KP', '050111040319'),
 ('Program', 'DIPLOMA SAINS KOMPUTER'), ('No. Matrik', ''),
 ('Nama Penyelia', 'NUR RASYIDAH'), ('Jawatan Penyelia', 'HR'),
 ('No. Tel (H) Penyelia', '017-9436975'), ('Alamat semasa', 'NO.3,JALAN'),
 ('Poskod', ''), ('Bandar', ''), ('Negeri', ''), ('No. Tel (R)', ''), ('No. Tel (H)', ''),
 ('Nama Ibu / Bapa / Penjaga', ''), ('Alamat Ibu / Bapa / Penjaga', ''), ('No. Tel Penjaga', ''),
 ('Nama & Alamat orang yang perlu dihubungi semasa kecemasan', ''), ('No. Tel Kecemasan', '')]
t = table(len(details), 2, [2.35, 4.45])
for row, (label, value) in zip(t.rows, details):
    row.cells[0].paragraphs[0].add_run(label).bold = True
    p = row.cells[1].paragraphs[0]
    if value: p.add_run(value)
    else: field(p, 'Klik di sini untuk isi.')
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))

newpage('PENULISAN BUKU LOG')
doc.add_paragraph(original(22))
doc.add_paragraph(original(25), 'Heading 2'); doc.add_paragraph(original(27))
doc.add_paragraph(original(30), 'Heading 2'); doc.add_paragraph(original(32))
import re
items = re.split(r'(?=\b(?:i|ii|iii|iv)\))', original(34))
for item in items:
    if item.strip(): doc.add_paragraph(item.strip())
doc.add_paragraph(original(35))
doc.add_paragraph(original(39), 'Heading 2')
for i in range(42, 53):
    p = doc.add_paragraph(original(i)); p.paragraph_format.space_after = Pt(3)

headings = ['PERANCANGAN TUGASAN YANG DIJALANKAN:', 'PERLAKSANAAN TUGASAN YANG DIJALANKAN:', 'HASIL TUGASAN YANG DIJALANKAN:', 'KESIMPULAN KESELURUHAN TUGASAN YANG DIJALANKAN:']
for week in range(1, 25):
    newpage(f'MINGGU {week}')
    p = doc.add_paragraph(); p.add_run('TARIKH: ').bold = True; field(p, 'Isi tarikh')
    p.add_run('     HARI: ').bold = True; field(p, 'Isi hari')
    for title in headings:
        doc.add_paragraph(title, 'Heading 2')
        t = table(1, 1, [6.8]); row = t.rows[0]
        row.height = Inches(.85); row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        field(row.cells[0].paragraphs[0])
    p = doc.add_paragraph('PENGESAHAN PENYELIA', 'Heading 2')
    t = table(1, 1, [6.8]); cell = t.cell(0,0)
    for j, label in enumerate(['Disemak oleh', 'Nama Pegawai / Penyelia Industri', 'Jabatan / Bahagian', 'Tarikh', 'Ulasan']):
        p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.add_run(label + ': ').bold = True
        if j == 0:
            p.paragraph_format.space_after = Pt(16)
        else: field(p, 'Klik di sini untuk isi.')
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))

p = section.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(); e = OxmlElement('w:fldSimple'); e.set(qn('w:instr'), 'PAGE'); r._r.addnext(e)
doc.core_properties.title = 'Buku Log Latihan Industri'
doc.core_properties.author = ''
path = OUT / 'Logbook_Template_Mudah_Isi.docx'
doc.save(path)
print(path)
print('weeks=24; fillable_fields=', field_id - 100)
